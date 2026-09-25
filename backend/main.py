"""FastAPI app. Endpoints:
GET  /api/options      dropdown values
POST /api/assistant    message + optional file + current form -> {intent, reply, fields, changed} (LangGraph)
POST /api/deviations   reviewed form    -> saved record
GET  /api/deviations   list saved records
"""
import io
from datetime import datetime
from typing import Optional

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from db import SessionLocal, Deviation
import json
from graph import run_assistant, SOURCES, IMPACTS, SEVERITIES

app = FastAPI(title="Deviation Intake")
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173"], allow_methods=["*"], allow_headers=["*"])

MAX_BYTES = 10 * 1024 * 1024


def file_to_text(name: str, data: bytes) -> str:
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
    if ext == "docx":
        from docx import Document
        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in r.cells) for r in t.rows]
        return "\n".join(parts)
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True)
        return "\n".join(" | ".join("" if c is None else str(c) for c in row)
                         for ws in wb for row in ws.iter_rows(values_only=True))
    if ext in ("txt", "csv", "eml"):
        return data.decode("utf-8", errors="ignore")
    raise HTTPException(400, f".{ext} is not supported yet. Upload PDF, DOCX, XLSX or TXT, or paste the text.")


@app.get("/api/options")
def options():
    return {"sources": SOURCES, "impacts": IMPACTS, "severities": SEVERITIES}


@app.post("/api/assistant")
async def assistant(message: str = Form(""), form: str = Form("{}"), force_extract: bool = Form(False),
                    file: Optional[UploadFile] = File(None)):
    """Single entry point for the AI panel: typed message, pasted text and/or an uploaded document."""
    doc_text = ""
    if file:
        data = await file.read()
        if len(data) > MAX_BYTES:
            raise HTTPException(400, "File is larger than 10 MB.")
        doc_text = file_to_text(file.filename, data)
        if len(doc_text.strip()) < 20:
            raise HTTPException(400, "No readable text found in the file. Try a text-based PDF or paste the text.")
    if len(message.strip()) < 2 and not doc_text:
        raise HTTPException(400, "Type a message, paste text or upload a document.")
    try:
        return run_assistant(message, doc_text, json.loads(form or "{}"), force_extract)
    except Exception as e:
        raise HTTPException(502, f"AI service error: {e}")


class DeviationIn(BaseModel):
    site: str
    date_of_occurrence: str
    title: str
    source: str
    product: Optional[str] = None
    batch_number: Optional[str] = None
    description: str
    impact: str
    severity: str
    ai_reason: Optional[str] = None


@app.post("/api/deviations")
def save(body: DeviationIn):
    if body.impact not in IMPACTS or body.severity not in SEVERITIES:
        raise HTTPException(422, "Invalid impact or severity.")
    with SessionLocal() as s:
        row = Deviation(**body.model_dump())
        s.add(row)
        s.flush()
        row.deviation_no = f"DEV-{datetime.utcnow().year}-{row.id:04d}"
        s.commit()
        s.refresh(row)
        return row.to_dict()


@app.get("/api/deviations")
def list_all():
    with SessionLocal() as s:
        return [r.to_dict() for r in s.query(Deviation).order_by(Deviation.id.desc()).limit(50)]
